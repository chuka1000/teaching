#!/usr/bin/env python3
"""
Expressions From Shapes And Prices (8CN, Lesson 8). Every expression, table and check value in the deck, the
worksheet and the speaker notes is derived here with sympy, and the strings typed in the files are READ BACK and
compared, so the files and the check cannot drift apart.

    python3 build/shapes-and-prices-check.py

Part 1 checks the worksheet data in build/shapes-and-prices-answers.js (the one module the worksheet and its
upside-down answers read). Part 2 reads the finished deck and worksheet. Part 2 is skipped if they are not built yet.
"""
import json, pathlib, re, subprocess, zipfile, datetime, difflib
from sympy import symbols, Eq, solve, expand, Poly, Rational as R
from sympy.parsing.sympy_parser import parse_expr, standard_transformations, implicit_multiplication_application, convert_xor

ROOT = pathlib.Path(__file__).parent.parent
LESSON = 'Expressions From Shapes And Prices'
OUT = ROOT / 'out' / LESSON
x, y, m, c, r, a, s = symbols('x y m c r a s')
TR = standard_transformations + (implicit_multiplication_application, convert_xor)
checks = 0


def ok(cond, msg):
    global checks
    assert cond, f'FAIL: {msg}'
    checks += 1
    print('ok:', msg)


def P(text):
    """Parse the deck's maths text (a real minus, ×, ÷) into sympy."""
    t = text.replace('−', '-').replace('×', '*').replace('÷', '/').replace('£', '')
    return parse_expr(t, transformations=TR, local_dict={'x': x, 'r': r, 'a': a, 's': s, 'm': m})


def terms(text):
    return len(re.split(r'\s[+−]\s', text))


def table_rule(xs, ys):
    sol = solve([Eq(m * xs[0] + c, ys[0]), Eq(m * xs[1] + c, ys[1])], (m, c))
    mm, cc = sol[m], sol[c]
    ok(all(mm * a_ + cc == b_ for a_, b_ in zip(xs, ys)), f'table x = {xs}, y = {ys} is linear: y = {mm}x {"+" if cc >= 0 else "-"} {abs(cc)}')
    return mm, cc


def show(mm, cc):
    return f'y = {mm}x {"+" if cc >= 0 else "−"} {abs(cc)}' if mm != 1 else f'y = x {"+" if cc >= 0 else "−"} {abs(cc)}'


# ------------------------------------------------------------------ the worksheet data ----
dump = subprocess.check_output(['node', '-e', "const a=require('./build/shapes-and-prices-answers.js');console.log(JSON.stringify({a:a.A,S:a.SHAPES,L:a.L_BOX,C:a.COST,G:a.AGES,T:a.T,R:a.RULE,U:a.STUDENT}))"], cwd=ROOT).decode()
D = json.loads(dump)
ANS, SHAPES, COST, AGES, TB, RULE, STUDENT = D['a'], D['S'], D['C'], D['G'], D['T'], D['R'], D['U']

for k, sh in SHAPES.items():
    sides = [P(sd) for sd in sh['sides']] if 'sides' in sh else [P(sh['side'])] * sh['n']
    total = expand(sum(sides))
    ok(total == P(sh['ans']), f'{k} ({sh["name"]}): the perimeter is {total}, the answer says {sh["ans"]}')
    ok(terms(sh['ans']) == 2 and sh['ans'].index('x') < re.search(r'\d+$', sh['ans']).start(), f'{k}: the answer {sh["ans"]} is collected (one x term, one number)')
    v = sh['chk']
    lengths = [sd.subs(x, v) for sd in sides]
    ok(all(l_ > 0 for l_ in lengths), f'{k}: at the check value x = {v} every side is positive: {lengths}')
    ok(sum(lengths) == P(sh['ans']).subs(x, v), f'{k}: the check x = {v} agrees: {sum(lengths)} = {P(sh["ans"]).subs(x, v)}')
    if sh['name'] == 'triangle':
        sl = sorted(lengths)
        ok(sl[0] + sl[1] > sl[2], f'{k}: at x = {v} the triangle {lengths} exists (two short sides are longer than the long one)')

# the L-shape's six sides really close up: bottom 3x + 2, right lower x + 2, step x, step 2, top 2x + 2, left x + 4
lsh = [P(sd) for sd in SHAPES['g15']['sides']]
bottom, right_low, step_h, step_v, top, left = lsh
ok(expand(left - right_low) == step_v and expand(bottom - top) == step_h, 'Gold Q15: the L-shape closes (left − right_low = 2 and bottom − top = x)')
box = [P(sd) for sd in D['L']['sides']]
ok(expand(sum(box)) == P(D['L']['ans']) == P(SHAPES['g15']['ans']), 'Gold Q15(b): the surrounding rectangle 3x + 2 by x + 4 has the same perimeter, 8x + 12')

for k, co in COST.items():
    tot = expand(sum(P(p_) for p_ in co['parts']))
    ok(tot == P(co['ans']), f'cost {k}: {co["parts"]} add to {tot}, the answer says {co["ans"]}')
    v = co['chk']; var = 's' if k == 's9' else 'x'
    sym = s if var == 's' else x
    ok(sum(P(p_).subs(sym, v) for p_ in co['parts']) == P(co['ans']).subs(sym, v), f'cost {k}: the check value {v} agrees ({sum(P(p_).subs(sym, v) for p_ in co["parts"])})')
for k, ag in AGES.items():
    tot = expand(sum(P(p_) for p_ in ag['parts']))
    var = r if k == 'b4' else a
    ok(tot == P(ag['ans']), f'ages {k}: {ag["parts"]} add to {tot}, the answer says {ag["ans"]}')
    v = ag['chk']
    ok(sum(P(p_).subs(var, v) for p_ in ag['parts']) == P(ag['ans']).subs(var, v), f'ages {k}: the check value {v} agrees')

for k, tb in TB.items():
    mm, cc = table_rule(tb['x'], tb['y'])
    ok(show(mm, cc) == RULE[k], f'table {k}: sympy finds "{show(mm, cc)}", the module says "{RULE[k]}"')
    ok(cc != tb['y'][0], f'table {k}: the number added ({cc}) is not the first y ({tb["y"][0]}), so the "first y" slip is wrong here')
ok(table_rule(TB['b5']['x'], TB['b5']['y']) == (2, 4) and table_rule(TB['tx']['x'], TB['tx']['y']) == (2, 6), 'taxi tables: £2 a km with £4 fixed (6, 8, 10) and £2 a km with £6 fixed (8, 10, 12)')
# the value of y at x = 0 is the number added: it is read off by going back, and it matches
for k in ('b5', 's11', 's13', 'g16', 'tx'):
    mm, cc = table_rule(TB[k]['x'], TB[k]['y'])
    back = TB[k]['y'][0] - mm * (TB[k]['x'][0] - 0)
    ok(back == cc, f'table {k}: going back to x = 0 gives y = {back}, the number added')
ok(R(27 - 9, 8 - 2) == 3 and 9 - 3 * 2 == 3, 'Silver Q11: x up 2, y up 6, so ×3; 9 − 6 = 3 at x = 0')
# the student slip
mm13, cc13 = table_rule(TB['s13']['x'], TB['s13']['y'])
ok(STUDENT['s13'] == 'y = 3x + 10' and 3 * 2 + 10 != TB['s13']['y'][1] and (mm13, cc13) == (3, 7), 'Silver Q13: the student\'s y = 3x + 10 fits the first pair? 3 + 10 = 13 is not 10, and fails the second (16, not 13); the rule is 3x + 7')
ok(3 * 1 + 10 != TB['s13']['y'][0], 'Silver Q13: the student used the first y (10) as the number added, so it does not even fit the first pair')
# Bronze Q6: the student's slip
sides6 = [P('x + 2'), 3, P('x + 2'), 3]
ok(expand(sum(sides6)) == P(STUDENT['b6ans']) and expand(sum(sides6)) != P(STUDENT['b6']) and expand(sum(sides6) - P(STUDENT['b6'])) == 3, 'Bronze Q6: the correct perimeter is 2x + 10, and the student\'s 2x + 7 is 3 short (one 3 missing)')
# Gold Q17: two taxi firms
ok(2 * 1 + 3 == 5 and 1 + 5 == 6 and 2 * 2 + 3 == 7 and 2 + 5 == 7 and 2 * 4 + 3 == 11 and 4 + 5 == 9, 'Gold Q17: A is 5 and B is 6 at 1 km; both 7 at 2 km; A is 11 and B is 9 at 4 km')
ok(solve(Eq(2 * x + 3, x + 5), x) == [2], 'Gold Q17: the two firms cost the same at exactly 2 km')
# Gold Q19
ok(expand(P('12x + 8') / 4) == P('3x + 2') and (12 + 8) / 4 == 5, 'Gold Q19: (12x + 8) ÷ 4 = 3x + 2, check x = 1: side 5, perimeter 20')

# every answer string that quotes an expression: each quoted result must be the right result
QUOTE = {'1': '8x + 4', '2': '8x + 12', '3': '3x + 50', '4': '2r + 3', '6': '2x + 10', '7': '6x + 6', '8': '7x + 2', '9': '5s − 4', '10': '4a + 25',
         '13': '3x + 7', '14': '5x + 5', '15': '8x + 12', '16': '8x + 15', '19': '3x + 2'}
for q_, e_ in QUOTE.items():
    ok(e_ in ANS[q_], f'answer {q_} says "{e_}"')
ok('y = 2x + 4' in ANS['5'] and 'y = 3x + 3' in ANS['11'] and 'y = 3x + 7' in ANS['13'], 'answers 5, 11 and 13 say the rules from the tables')
ok(expand(2 * P('2x + 3') + P('x - 1')) == P('5x + 5') and expand(2 * P('3x + 2') + 2 * P('x + 4')) == P('8x + 12') and expand(P('3*s') + 2 * (s - 2)) == P('5s - 4'), 'bracket expansions quoted in the answers: 2(2x + 3) + (x − 1), 2(3x + 2) + 2(x + 4), 3s + 2(s − 2)')
ok(len(ANS) == 19 and all(ANS[str(i)] for i in range(1, 20)), 'nineteen worksheet answers, one for every question')

# the "x = 0" idea is right in every table the lesson uses
for xs_, ys_ in ((TB['tx']['x'], TB['tx']['y']), (TB['b5']['x'], TB['b5']['y'])):
    mm, cc = table_rule(xs_, ys_)
    ok(ys_[0] - mm == cc, f'going back one step from x = {xs_[0]} (y = {ys_[0]}) takes away {mm} and gives {cc}')

# ---- the Do Now, Cold Call, We Do, Hook, I Dos, Plenary: the numbers typed in the deck ----
N = {}
ok(expand(2 * (x + 6) - 3 * (x - 1)) == -x + 15 and 2 * 8 - 3 * 1 == 13 == -2 + 15, 'Do Now Q1: 2(x + 6) − 3(x − 1) = 2x + 12 − 3x + 3 = −x + 15 (check x = 2: 13 and 13)')
ok(expand(8 - (x + 3)) == 5 - x and expand(8 - (x + 3)) != 8 - x + 3, 'Do Now Q2: 8 − (x + 3) = 8 − x − 3 = 5 − x (the student\'s 8 − x + 3 is 11 − x)')
mm_, cc_ = table_rule([1, 2, 3], [12, 17, 22])
ok((mm_, cc_) == (5, 7) and cc_ != 12, 'Do Now Q3: x = 1, 2, 3 and y = 12, 17, 22 gives y = 5x + 7')
ok(4 * 5 - 3 * 2 == 14, 'Do Now Q4: 4x − 3y at x = 5, y = 2 is 20 − 6 = 14')
ok(9 + 4 + 9 + 4 == 26 and 2 * (9 + 4) == 26, 'Do Now Q5: the perimeter of a 9 cm by 4 cm rectangle is 26 cm')
mt, ct = table_rule([1, 2, 3], [8, 10, 12])
ok((mt, ct) == (2, 6) and 8 - 2 == 6, 'Do Now Q6: the taxi (8, 10, 12) is y = 2x + 6, so a 0 km journey costs £6')

hook = [P('3x + 3'), P('6x + 6'), P('4x + 6')]
true_perim = expand((x + 3) + 2 * x + (x + 3) + 2 * x)
ok(true_perim == hook[1] and expand((x + 3) + 2 * x) == hook[0] and expand(2 * (x + 3) + 2 * x) == hook[2], 'Hook: x + 3 and 2x gives 6x + 6 (B); A = 3x + 3 adds two sides only; C = 4x + 6 doubles the x + 3 but counts 2x once')
ok(expand(2 * (x + 3) + 2 * (2 * x)) == 6 * x + 6 and 2 * 7 + 2 * 8 == 30 == 6 * 4 + 6, 'I Do 1: both routes give 6x + 6 and the check x = 4 (sides 7 and 8) gives 30')
ok([expand(t_) for t_ in (x + 3, (x + 3) + 2 * x, (x + 3) + 2 * x + (x + 3), (x + 3) + 2 * x + (x + 3) + 2 * x)] == [x + 3, 3 * x + 3, 4 * x + 6, 6 * x + 6], 'the running total in the animation is x + 3, 3x + 3, 4x + 6, 6x + 6')
ok(4 * 10 + 30 == 70 and expand(4 * x + 30) == 4 * x + 30, 'I Do 2 cost: 4 pens at x pence and a 30p ruler is 4x + 30 (x = 10: 70p)')
ok(expand(m + (m + 4) + 3 * m) == 5 * m + 4 and 12 + 16 + 36 == 64 == 5 * 12 + 4, 'I Do 2 age: m + (m + 4) + 3m = 5m + 4 (m = 12: 12 + 16 + 36 = 64)')
ok(expand((x + 2) + (x + 2) + (2 * x - 1)) == 4 * x + 3 and 7 + 7 + 9 == 23 == 4 * 5 + 3, 'We Do row 1: (x + 2) + (x + 2) + (2x − 1) = 4x + 3 (x = 5: 7 + 7 + 9 = 23)')
ok(expand(s + (s - 2) + 4 * s) == 6 * s - 2 and 10 + 8 + 40 == 58 == 6 * 10 - 2, 'We Do row 2: s + (s − 2) + 4s = 6s − 2 (s = 10: 10 + 8 + 40 = 58)')
ok(2 * 3 + 6 == 12, 'We Do row 3: y = 2x + 6 at x = 3 is 12 (the last y)')

# Cold Call
ok(expand(2 * (x + 5) + 2 * x) == 4 * x + 10, 'Cold Call Q1: x + 5 and x gives 4x + 10')
ok(expand(2 * x + (2 * x + 1) + (x - 4)) == 5 * x - 3, 'Cold Call Q2: 2x, 2x + 1 and x − 4 gives 5x − 3')
ok(expand(5 * x + 40) == 5 * x + 40 and 5 * 10 + 40 == 90, 'Cold Call Q3: 5 buns at x pence and a 40p drink cost 5x + 40')
mm_, cc_ = table_rule([1, 2, 3], [9, 15, 21])
ok((mm_, cc_) == (6, 3) and 9 - 6 == 3, 'Cold Call Q4: x = 1, 2, 3 and y = 9, 15, 21 gives y = 6x + 3, and y = 3 at x = 0')
ok(expand(-6 * (2 * x - 1)) == -12 * x + 6, 'Cold Call Q5: −6(2x − 1) = −12x + 6')
ok(expand(5 * (x + 2) + 3 * (x - 4)) == 8 * x - 2, 'Cold Call Q6: 5(x + 2) + 3(x − 4) = 5x + 10 + 3x − 12 = 8x − 2')

# Plenary
ok(expand((x + 3) + 2 * x) != true_perim, 'Plenary Q1 is FALSE: x + 3 and 2x have perimeter 6x + 6, not 3x + 3')
ok(expand(4 * x + 30) == 4 * x + 30, 'Plenary Q2 is TRUE: 4x + 30')
ok(expand(m + 3 * m) == 4 * m and 4 * m != 3 * m + 1, 'Plenary Q3 is FALSE: m + 3m = 4m, not 3m + 1')
ok(table_rule([1, 2, 3], [8, 10, 12]) == (2, 6) and (2, 8) != (2, 6), 'Plenary Q4 is FALSE: the table gives 2x + 6, not 2x + 8')
ok(2 * 0 + 6 == 6, 'Plenary Q5 is TRUE: in y = 2x + 6 the 6 is y when x = 0')

# ---------------------------------------------------------------- part 2: the finished files ----
DECK = OUT / f'{LESSON}.pptx'
SHEET = OUT / f'{LESSON} worksheet.docx'
if DECK.exists():
    from pptx import Presentation
    prs = Presentation(str(DECK)); slides = list(prs.slides)

    def named(i):
        out = {}
        def go(shapes):
            for sh in shapes:
                if sh.shape_type == 6: go(sh.shapes)
                elif sh.has_text_frame and sh.text_frame.text.strip(): out[sh.name] = sh.text_frame.text
        go(slides[i - 1].shapes); return out
    def text(i): return ' | '.join(v.replace('\n', ' ') for v in named(i).values())
    def notes(i): return slides[i - 1].notes_slide.notes_text_frame.text

    pills = []
    for i in range(1, len(slides) + 1):
        mm = re.search(r'([A-Z][A-Z ]+) · (\d+) MIN', text(i)); pills.append((mm.group(1), int(mm.group(2))) if mm else None)
    ok(len(slides) == 10 and all(pills), 'ten slides, each with a phase pill')
    ok([p_[0] for p_ in pills] == ['DO NOW', 'TODAY', 'HOOK', 'I DO', 'I DO', 'WE DO', 'COLD CALL', 'YOU DO', 'MARK', 'PLENARY'], 'phases in the TEMPLATE.md order')
    ok([p_[1] for p_ in pills] == [10, 1, 2, 3, 3, 5, 6, 14, 3, 3] and sum(p_[1] for p_ in pills) == 50, 'minutes 10, 1, 2, 3, 3, 5, 6, 14, 3, 3 add to 50')
    ok(all(notes(i).strip() for i in range(1, 11)), 'speaker notes on every slide')
    ok(prs.core_properties.title == LESSON and prs.core_properties.subject == 'Y8 Maths · Algebra · Lesson 8 · 8CN', f'dc:title and dc:subject are set, with the class: {prs.core_properties.subject}')
    d = re.search(r'(Monday|Tuesday|Wednesday|Thursday|Friday|Saturday|Sunday) (\d+) (\w+) (\d{4})', text(1))
    dt = datetime.datetime.strptime(f'{d.group(2)} {d.group(3)} {d.group(4)}', '%d %B %Y')
    ok(dt.strftime('%A') == d.group(1) and d.group(0) == 'Tuesday 20 October 2026', f'the date on slide 1 is a real {d.group(1)}: {d.group(0)}')
    ok(prs.slide_width / 914400 == 10 and abs(prs.slide_height / 914400 - 5.625) < 1e-6, 'maths canvas: 10 x 5.625')
    ok('Objectives' in named(2)['slide_title'], 'slide 2 is titled Objectives')

    # the Do Now answers, parsed from the slide and compared with sympy
    n1 = named(1)
    ok(P(n1['d0_a'].split('=')[-1]) == -x + 15 and n1['d0_a'].endswith('= −x + 15'), f'Do Now Q1 on the slide: "{n1["d0_a"]}"')
    ok('5 − x' in n1['d1_a'] and '8 − x + 3' in n1['d1_q'], f'Do Now Q2 on the slide: "{n1["d1_a"]}"')
    ok('y = 5x + 7' in n1['d2_a'] and 'y = 12, 17, 22' in n1['d2_q'], f'Do Now Q3 on the slide: "{n1["d2_a"]}"')
    ok(n1['d3_a'].startswith('4 × 5 − 3 × 2 = 14'), f'Do Now Q4 on the slide: "{n1["d3_a"]}"')
    ok(n1['d4_a'].startswith('9 + 4 + 9 + 4 = 26 cm'), f'Do Now Q5 on the slide: "{n1["d4_a"]}"')
    ok(n1['d5_a'].startswith('£6') and 'Suggest' in n1['d5_q'], f'Do Now Q6 (the preview) on the slide: "{n1["d5_a"]}"')

    # no Do Now question repeats one from the last three lessons in the class
    prev = ['Function Machines', 'Putting Numbers In', 'Expand and Simplify', 'Brackets and Beyond']
    old = []
    for name in prev:
        pr = Presentation(str(ROOT / 'reference' / f'{name}.pptx'))
        for sl in ([pr.slides[0]] + ([pr.slides[6]] if len(pr.slides) > 6 else [])):
            for sh in sl.shapes:
                if sh.has_text_frame and len(sh.text_frame.text.strip()) > 12: old.append(sh.text_frame.text.replace('\n', ' '))
    mine = [n1[f'd{i}_q'] for i in range(6)]
    worst = max((difflib.SequenceMatcher(None, a_.lower(), b_.lower()).ratio(), a_, b_) for a_ in mine for b_ in old)
    print('closest pair by wording:', worst[1][:60], '<>', worst[2][:60])
    digits = lambda s_: re.findall(r'\d+', s_)
    repeats = [(a_, b_) for a_ in mine for b_ in old if digits(a_) and digits(a_) == digits(b_) and difflib.SequenceMatcher(None, a_.lower(), b_.lower()).ratio() > 0.5]
    ok(not repeats, f'no Do Now question repeats one from the last lessons ({", ".join(prev)}): same skills, new numbers (closest by wording {worst[0]:.2f}); repeats: {repeats}')

    # Hook
    n3 = named(3)
    ok([n3['h0_t'], n3['h1_t'], n3['h2_t']] == ['3x + 3', '6x + 6', '4x + 6'], 'Hook options on the slide: 3x + 3, 6x + 6, 4x + 6')
    # I Do 1 and 2
    n4 = named(4)
    ok('(x + 3) + 2x + (x + 3) + 2x' in n4['ea_l1'] and n4['ea_l2'].endswith('6x + 6') and '2(x + 3) + 2(2x)' in n4['ea_l4'] and n4['ea_l5'].endswith('6x + 6'), 'I Do 1 on the slide: both routes end at 6x + 6')
    n5 = named(5)
    ok(n5['co_l2'].endswith('4x + 30') and n5['ag_l3'].endswith('5m + 4'), f'I Do 2 on the slide: cost ends "{n5["co_l2"]}", age ends "{n5["ag_l3"]}"')
    # We Do
    n6 = named(6)
    ok('4x + 3' in (n6.get('wf0_2_a') or n6.get('wf0_2_q')) and '6s − 2' in (n6.get('wf1_2_a') or '') and 'y = 2x + 6' in (n6.get('wf2_2_a') or ''), 'We Do on the slide: 4x + 3, 6s − 2 and y = 2x + 6')
    # Cold Call
    n7 = named(7)
    CC = {0: '4x + 10', 1: '5x − 3', 2: '5x + 40', 3: 'y = 6x + 3', 4: '−12x + 6', 5: '8x − 2'}
    for i, e_ in CC.items(): ok(e_ in n7[f'c{i}_a'], f'Cold Call Q{i + 1} on the slide says "{e_}"')
    # Plenary
    n10 = named(10)
    truth = [False, True, False, False, True]
    ok([n10[f'p{i}_v'] for i in range(5)] == ['TRUE' if t_ else 'FALSE' for t_ in truth], f'Plenary labels agree with the maths: {[n10[f"p{i}_v"] for i in range(5)]}')

if SHEET.exists():
    z = zipfile.ZipFile(SHEET)
    body = re.sub(r'\s+', ' ', re.sub(r'<[^>]+>', ' ', z.read('word/document.xml').decode('utf-8')))
    for k, sh in SHAPES.items():
        if 'sides' in sh and k not in ('g15',):
            ok(all(sd in body for sd in sh['sides']) or all(sd.replace('−', '-') in body for sd in sh['sides']), f'the worksheet prints the sides of {k}: {sh["sides"]}')
    BLANK = {'b5', 'g16', 'g17a', 'g17b'}          # tables the student completes: only x is printed
    for k, tb in TB.items():
        if k == 'tx': continue
        vals = tb['x'] if k in BLANK else tb['x'] + tb['y']
        ok(all(re.search(rf'\b{v}\b', body) for v in vals), f'the worksheet prints table {k}: x = {tb["x"]}' + ('' if k in BLANK else f', y = {tb["y"]}')) 
    src = (ROOT / 'build' / 'shapes-and-prices-worksheet.js').read_text(encoding='utf-8')
    for i in range(1, 20): ok(re.search(rf"q\('{i}',", src), f'worksheet has question {i}')
    ok('Why does that step work?' in src and "q('12'" in src, 'the "why does that step work?" prompt is Q12, on the x = 0 step')
    ok('Where would you meet this' in src and "q('18'" in src, 'the "where would you meet this outside the lesson?" question is Q18')

print(f'\n{checks} checks passed')
