#!/usr/bin/env python3
"""
Checks everything numeric (and every claim that can be checked) in Shuffling The Gene Pool (8I): the deck, the bead activity, the
animation's model, the worksheet and its answers, and that the game the deck names exists. Values are read BACK OUT of the finished
files and recomputed here with exact fractions (sympy for the binomial figures behind the We Do), independently of the JavaScript and
the Python that wrote them. The game is checked over 300 generated games by build/check-shuffling-the-gene-pool-game.py.

    python3 build/shuffling-the-gene-pool-check.py
"""
import datetime, difflib, json, pathlib, re, subprocess, zipfile
from fractions import Fraction as F
from sympy import Rational as R, binomial
from pptx import Presentation

ROOT = pathlib.Path(__file__).parent.parent
LESSON = 'Shuffling The Gene Pool'
OUT = ROOT / 'out' / LESSON
DECK = OUT / f'{LESSON}.pptx'
checks = 0
def ok(cond, msg):
    global checks
    assert cond, f'FAIL: {msg}'
    checks += 1
    print('ok:', msg)
pct = lambda k, n: F(100 * k, n)

prs = Presentation(str(DECK)); slides = list(prs.slides)
def named(i):
    out = {}
    def go(shapes):
        for sh in shapes:
            if sh.shape_type == 6: go(sh.shapes)
            elif sh.has_text_frame and sh.text_frame.text.strip(): out[sh.name] = sh.text_frame.text
    go(slides[i - 1].shapes); return out
def text(i): return ' | '.join(v.replace('\n', ' ') for v in named(i).values())
def notes(i): return slides[i - 1].notes_slide.notes_text_frame.text

# ---- the shape: TEMPLATE.md ----
pills = []
for i in range(1, len(slides) + 1):
    m = re.search(r'([A-Z][A-Z ]+) · (\d+) MIN', text(i)); pills.append((m.group(1), int(m.group(2))) if m else None)
ok(len(slides) == 10 and all(pills), 'ten slides, each with a phase pill')
ok([p[0] for p in pills] == ['DO NOW', 'TODAY', 'HOOK', 'I DO', 'I DO', 'WE DO', 'COLD CALL', 'YOU DO', 'MARK', 'PLENARY'], 'phases in the TEMPLATE.md order')
ok([p[1] for p in pills] == [10, 1, 2, 3, 3, 5, 6, 14, 3, 3] and sum(p[1] for p in pills) == 50, 'minutes 10, 1, 2, 3, 3, 5, 6, 14, 3, 3 add to 50')
ok(all(notes(i).strip() for i in range(1, 11)), 'speaker notes on every slide')
ok(prs.core_properties.title == LESSON and prs.core_properties.subject == 'Y8 Science · Natural Selection · Lesson 7 · 8I', f'dc:title and dc:subject are set, with the class: {prs.core_properties.subject}')
d = re.search(r'(Monday|Tuesday|Wednesday|Thursday|Friday|Saturday|Sunday) (\d+) (\w+) (\d{4})', text(1))
dt = datetime.datetime.strptime(f'{d.group(2)} {d.group(3)} {d.group(4)}', '%d %B %Y')
ok(dt.strftime('%A') == d.group(1) and dt.date() == datetime.date(2026, 10, 7), f'the date on slide 1 is a real {d.group(1)}, and it is the day after the build (the brief says "tomorrow"): {d.group(0)}')
n2 = named(2)
OBJ = ['Identify four forces that change allele frequencies.', 'Describe genetic drift as a random change.', 'Explain why genetic drift has more effect in small populations.']
ok(all(n2[f'o{i}_t'] == OBJ[i] for i in range(3)) and 'Objectives' in n2['slide_title'] and text(2).startswith('TODAY · 1 MIN'), 'the three objectives are exactly the brief\'s wording, under the title Objectives with the TODAY pill')

# ---- Do Now ----
n1 = named(1)
ok(30 * 2 == 60 and pct(24, 60) == 40 and pct(24, 30) == 80 and n1['d0_a'].startswith('40%.'), 'Do Now Q1: 30 beetles × 2 = 60 alleles, 24 ÷ 60 = 40% (the trap, 24 ÷ 30, gives 80%)')
ok('5' in n1['d4_a'] and F(10, 2) == 5, 'Do Now Q5: a fair coin flipped 10 times gives 5 heads on average')
ok('other than natural selection' in n1['d5_q'] and 'all four forces' in n1['d5_a'], 'Do Now Q6 previews today: a force other than natural selection')
prev = ['Small Changes, Big Changes', 'More Evidence', 'Fossils And The Fossil Record']
def walk(shapes):
    for sh in shapes:
        if sh.shape_type == 6: yield from walk(sh.shapes)
        elif sh.has_text_frame and sh.text_frame.text.strip(): yield sh.text_frame.text.replace('\n', ' ')
old = []
for name in prev:
    f = ROOT / 'reference' / f'{name}.pptx'
    if not f.exists(): f = ROOT / 'out' / name / f'{name}.pptx'
    old += [t_ for t_ in walk(Presentation(str(f)).slides[0].shapes) if len(t_) > 12]
worst = max((difflib.SequenceMatcher(None, n1[f'd{i}_q'].lower(), b.lower()).ratio(), n1[f'd{i}_q'], b) for i in range(6) for b in old)
print('closest pair:', worst[1][:70], '<>', worst[2][:70])
ok(worst[0] < 0.7, f'no Do Now question repeats one from {", ".join(prev)}: closest match {worst[0]:.2f}')

# ---- Hook (3) ----
n3 = named(3)
ok(pct(10, 20) == 50 and '10 green and 10 brown' in n3['slide_title'] and 'C' == n3['h2_k'] and 'Probably not 50%' in n3['h2_t'], 'Hook: 10 green and 10 brown is 50%, and the answer is C (probably not 50%, which way nobody can say)')

# ---- I Do 1 (4) and I Do 2 (5) ----
n4, n5 = named(4), named(5)
ok(all(n4[k] for k in ('mut_h', 'gf_h', 'ns_h', 'gd_h')) and [n4[k] for k in ('mut_h', 'gf_h', 'ns_h', 'gd_h')] == ['MUTATION', 'GENE FLOW', 'NATURAL SELECTION', 'GENETIC DRIFT'], 'I Do 1 names the four forces')
ok('NOT RANDOM' in n4['ns_tag'] and n4['gd_tag'].startswith('RANDOM'), 'I Do 1: natural selection is NOT RANDOM, genetic drift is RANDOM (the key distinction)')
ok(4 * 2 == 8 and pct(4, 8) == 50 and pct(0, 2) == 0 and '4 beetles = 8 alleles' in n5['wx_t'] and '4 ÷ 8 = 50%' in n5['wx_t'] and '0 ÷ 2 = 0%' in n5['wx_t'], 'I Do 2 worked example: 4 beetles are 8 alleles, 4 ÷ 8 = 50%; the survivor with 2 brown alleles leaves 0 ÷ 2 = 0%')
ok('about 20 animals by 1892' in n5['sl_t'] and 'over 200,000' in n5['sl_t'], 'I Do 2: northern elephant seals, about 20 animals by 1892, now over 200,000')
ok('about 20 northern elephant seals' in notes(5) or 'only about 20 left' in notes(5), 'the I Do 2 notes give the seal facts')

# ---- the animation's model ----
sim = json.load((ROOT / 'build' / 'shuffling-the-gene-pool.sim.json').open())
G = sim['generations']
for side in ('small', 'big'):
    n = sim[side]['n']
    for path in sim[side]['runs']:
        ok(len(path) == G + 1 and path[0] == n // 2 and all(0 <= k <= n for k in path), f'model {side}: a run of {G} generations starts at half ({n // 2} of {n})')
        ok(all((path[i] not in (0, n)) or path[i + 1] == path[i] for i in range(len(path) - 1)), f'model {side}: once an allele reaches 0 or all of the pool it stays there (it is lost for good)')
def spread(side): return sum(abs(F(r[-1], sim[side]['n']) - F(1, 2)) for r in sim[side]['runs']) / len(sim[side]['runs']) * 100
ok(sim['small']['n'] == 4 and sim['big']['n'] == 200 and len(sim['small']['runs']) == len(sim['big']['runs']) == 6, 'model: 6 runs each of 4 alleles and 200 alleles')
ok(spread('small') > 2 * spread('big'), f'model: the small population ends {float(spread("small")):.1f} points from 50% on average, the big one {float(spread("big")):.1f}: the small one moves much further')
ok('Small population: {N_SMALL} alleles' in (ROOT / 'build' / 'media' / 'shuffling-the-gene-pool-media.py').read_text() and 'A computer model' in (ROOT / 'build' / 'media' / 'shuffling-the-gene-pool-media.py').read_text(), 'the animation says it is a computer model')

# ---- We Do (6): the bead activity, exactly (sympy) ----
n6 = named(6)
def stats(n):
    mean = sum(binomial(n, k) * R(abs(2 * k - n), 2 * n) * 100 for k in range(n + 1)) / 2 ** n
    p25 = sum(binomial(n, k) for k in range(n + 1) if R(abs(2 * k - n), 2 * n) * 100 >= 25) / R(2 ** n)
    return mean, p25
m4, p4 = stats(4); m20, p20 = stats(20)
ok(m4 == R(75, 4) and p4 == R(5, 8), f'a population of 4 alleles drawn from a 50:50 stock: mean distance from 50% is {float(m4)} points, and {p4} of draws are at least 25 points away')
ok(abs(float(m20) - 8.81) < 0.01 and abs(float(p20) - 0.0414) < 0.0001 and 23 < 1 / float(p20) < 25, f'a population of 20 alleles: mean distance {float(m20):.2f} points, and {float(p20):.4f} of draws (1 in {1 / float(p20):.1f}) are at least 25 points away')
ok('about 19 points' in n6['ex_t'] and '5 draws in 8' in n6['ex_t'] and 'about 9 points' in n6['ex_t'] and '1 draw in 24' in n6['ex_t'], 'the We Do slide says: about 19 points and 5 draws in 8 (small), about 9 points and 1 draw in 24 (big)')
ok(round(float(m4)) == 19 and round(float(m20)) == 9 and float(m4) / float(m20) > 2, 'those figures round correctly, and the small population moves more than twice as far')
ok('4 beads' in n6['sm_t'] and '20 beads' in n6['bg_t'] and 'of 4' in n6['sm_t'] and 'of 20' in n6['bg_t'], 'the procedure draws 4 beads (small) and 20 beads (big)')
ok('draw 10 beads at random from a big bag and from a small one' in notes(6) and 'THE CHANGE, AND WHY' in notes(6) and 'depends on HOW MANY beads are drawn' in notes(6), 'the notes explain why the brief\'s "10 beads from each bag" had to change')
ok(F(6, 16) == F(3, 8) and 'about 3 times in 8' in notes(6), 'the notes: exactly 2 green of 4 happens 6 times in 16, about 3 in 8')

# ---- Cold Call (7) and Plenary (10) ----
n7 = named(7)
ok(pct(15, 60) == 25 and n7['c4_a'].startswith('25%.'), 'Cold Call Q5: 15 of 60 alleles is 25%')
ok('Natural selection. Not random' in n7['c1_a'] and 'Genetic drift. Random' in n7['c2_a'], 'Cold Call Q2 and Q3 test the key distinction: natural selection is not random, drift is')
ok('Small Changes' in notes(7), 'the Cold Call notes say Q5 and Q6 are from the previous lesson')
n10 = named(10)
ok([n10[f'p{i}_v'] for i in range(5)] == ['FALSE', 'TRUE', 'FALSE', 'FALSE', 'TRUE'], 'Plenary: FALSE, TRUE, FALSE, FALSE, TRUE (each FALSE is a misconception from the lesson; Q3 is the applied item)')

# ---- the worksheet and its answers ----
js = subprocess.check_output(['node', '-e', "const a=require('./build/shuffling-the-gene-pool-answers.js');console.log(JSON.stringify({a:a,D:a.D}))"], cwd=ROOT).decode()
data = json.loads(js); A = {k: v for k, v in data['a']}; D = data['D']
ok(len(A) == 14 and all(A[str(i)] for i in range(1, 15)), 'fourteen worksheet answers, one for every question, in order')
ok(pct(*D['q5']['before']) == 50 and pct(*D['q5']['after']) == 100 and '= 50%' in A['5'] and '= 100%' in A['5'], 'Q5: 5 of 10 is 50%, then 4 of 4 is 100%')
ok(pct(*D['w']['small']) == 75 and abs(pct(*D['w']['small']) - 50) == 25 and pct(*D['w']['big']) == 60 and abs(pct(*D['w']['big']) - 50) == 10, 'Silver worked example: 3 of 4 is 75% (25 points); 12 of 20 is 60% (10 points)')
ok(pct(D['q6']['green'], D['q6']['n']) == 25 and abs(pct(D['q6']['green'], D['q6']['n']) - 50) == 25 and '= 25%' in A['6'] and '= 25 points' in A['6'], 'Q6: 1 of 4 is 25%, 25 points from 50%')
ok(pct(D['q7']['green'], D['q7']['n']) == 65 and abs(pct(D['q7']['green'], D['q7']['n']) - 50) == 15 and '= 65%' in A['7'] and '= 15 points' in A['7'], 'Q7: 13 of 20 is 65%, 15 points from 50%')
left, gl = D['q13']['n'] - D['q13']['lostAlleles'], D['q13']['green'] - D['q13']['lostGreen']
ok((left, gl) == (10, 2) and pct(gl, left) == 20 and '= 20%' in A['13'] and '30 percentage points' in A['13'], 'Q13: (20 − 10) alleles left, (10 − 8) = 2 green: 2 of 10 is 20%, a fall of 30 points')
ok(pct(D['q13']['green'], D['q13']['n']) == 50 and pct(D['q13']['green'] - D['q13']['lostGreen'], D['q13']['n']) != pct(gl, left), 'Q13: the shrinking pool matters (2 of 20 would be 10%, but 2 of 10 is 20%)')
z = zipfile.ZipFile(OUT / f'{LESSON} worksheet.docx')
body = re.sub(r'\s+', ' ', re.sub(r'<[^>]+>', ' ', z.read('word/document.xml').decode('utf-8')))
for needle in ['3 ÷ 4 = 0.75 = 75%', '12 ÷ 20 = 0.6 = 60%', '1 green', '13 green', '20 alleles', 'northern elephant seals']:
    ok(needle in body, f'the worksheet prints "{needle}"')
wsrc = (ROOT / 'build' / 'shuffling-the-gene-pool-worksheet.js').read_text()
for i in range(1, 15): ok(re.search(rf"q\('{i}',", wsrc), f'worksheet has question {i}')
ok("q('8'" in wsrc and 'Why does that step work?' in wsrc and 'percentage before comparing' in wsrc, 'the "why does that step work?" prompt (Q8) is on turning counts into percentages before comparing')
ok("q('12'" in wsrc and 'Where would you meet this idea outside the lesson' in wsrc, 'the "where would you meet this outside the lesson?" question is Q12')
ok((OUT / f'{LESSON} worksheet.docx').exists(), 'the worksheet exists')

# ---- the game: the deck points at a file, so the file must exist ----
ok(f'{LESSON} game' in text(8), 'slide 8 names "Shuffling The Gene Pool game"')
game = OUT / f'{LESSON} game.html'
ok(game.exists() and game.stat().st_size > 20000, f'the game file the slide points at exists ({game.stat().st_size // 1024} KB)')
html = game.read_text(encoding='utf-8')
ok(html == (ROOT / 'build' / 'shuffling-the-gene-pool-game.html').read_text(encoding='utf-8'), 'the delivered game is exactly the source file')
ok(all(f in html for f in ('Mutation', 'Gene flow', 'Natural selection', 'Genetic drift')), 'the game has the four forces')
print(f'\n{checks} checks passed')
