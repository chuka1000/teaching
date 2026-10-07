#!/usr/bin/env python3
"""
Checks everything numeric (and every claim that can be checked) in When The Resultant Is Zero (10A): the deck, the animation's numbers, the
worksheet and its answers, the free-body diagrams the worksheet prints, that "no resultant force" is never written as "no forces" where it
matters (the brief's AVOID), and that the game the deck names exists. Values are read BACK OUT of the finished files and recomputed here
(sympy, exact), independently of the JavaScript that wrote them. The game is checked over 300 generated games by
build/check-when-the-resultant-is-zero-game.py.

    python3 build/when-the-resultant-is-zero-check.py
"""
import ast, datetime, difflib, json, pathlib, re, subprocess, zipfile
from sympy import Rational as R, sympify
from pptx import Presentation

ROOT = pathlib.Path(__file__).parent.parent
LESSON = 'When The Resultant Is Zero'
OUT = ROOT / 'out' / LESSON
DECK = OUT / f'{LESSON}.pptx'
checks = 0
def ok(cond, msg):
    global checks
    assert cond, f'FAIL: {msg}'
    checks += 1
    print('ok:', msg)
G = R(98, 10)

prs = Presentation(str(DECK)); slides = list(prs.slides)
def named(i):
    out = {}
    def go(shapes):
        for sh in shapes:
            if sh.shape_type == 6: go(sh.shapes)
            elif sh.has_text_frame and sh.text_frame.text.strip(): out[sh.name] = sh.text_frame.text
    go(slides[i - 1].shapes); return out
def text(i): return ' | '.join(v.replace('\n', ' ') for v in named(i).values())
def notes(i): return slides[i - 1].notes_slide.notes_text_frame.text

# ---- the shape: TEMPLATE.md ----
pills = []
for i in range(1, len(slides) + 1):
    m = re.search(r'([A-Z][A-Z ]+) · (\d+) MIN', text(i)); pills.append((m.group(1), int(m.group(2))) if m else None)
ok(len(slides) == 10 and all(pills), 'ten slides, each with a phase pill')
ok([p[0] for p in pills] == ['DO NOW', 'TODAY', 'HOOK', 'I DO', 'I DO', 'WE DO', 'COLD CALL', 'YOU DO', 'MARK', 'PLENARY'], 'phases in the TEMPLATE.md order')
ok([p[1] for p in pills] == [10, 1, 2, 3, 3, 5, 6, 14, 3, 3] and sum(p[1] for p in pills) == 50, 'minutes 10, 1, 2, 3, 3, 5, 6, 14, 3, 3 add to 50')
ok(all(notes(i).strip() for i in range(1, 11)), 'speaker notes on every slide')
ok(prs.core_properties.title == LESSON and prs.core_properties.subject == 'Y10 Science · Forces · 10A', f'dc:title and dc:subject are set, with the class: {prs.core_properties.subject}')
d = re.search(r'(Monday|Tuesday|Wednesday|Thursday|Friday|Saturday|Sunday) (\d+) (\w+) (\d{4})', text(1))
dt = datetime.datetime.strptime(f'{d.group(2)} {d.group(3)} {d.group(4)}', '%d %B %Y')
ok(dt.strftime('%A') == d.group(1) and dt.date() == datetime.date(2026, 10, 6), f'the date on slide 1 is a real {d.group(1)}, and it is the build date (the brief says "today"): {d.group(0)}')
n2 = named(2)
OBJ = ['State Newton’s first law. (P1.5.1.6)', 'Explain why an object at constant speed has no resultant force.', 'Explain what happens to an object with no forces acting at all.']
ok(all(n2[f'o{i}_t'] == OBJ[i] for i in range(3)) and 'Objectives' in n2['slide_title'] and text(2).startswith('TODAY · 1 MIN'), "the three objectives are exactly the brief's wording (objective 1 with its code), under the title Objectives with the TODAY pill")
syl = subprocess.run(['python3', 'tools/syllabus.py', 'P1.5.1.6'], cwd=ROOT, capture_output=True, text=True).stdout.lower()
ok('constant' in syl and 'resultant force' in syl and 'remains at rest' in syl, 'the scheme of work for P1.5.1.6 is the first law: at rest, or constant speed in a straight line, unless a resultant force')

# ---- Do Now ----
n1 = named(1)
ok(80 - 35 == 45 and n1['d0_a'].startswith('45 N to the right'), 'Do Now Q1: 80 − 35 = 45 N to the right')
ok(3 * 5 == 15 and n1['d2_a'].startswith('15 m/s'), 'Do Now Q3: v = 0 + 3 × 5 = 15 m/s')
ok('gradient' in n1['d3_q'] and n1['d3_a'].startswith('The speed'), 'Do Now Q4: the gradient of a distance-time graph is the speed')
ok('stops' in n1['d5_q'] and 'friction' in n1['d5_a'], 'Do Now Q6 previews today: a ball rolls and stops (friction)')
prev = ['Resultant Forces', 'Gravitational Fields and Free Fall', 'Weight and Gravity', 'Mass and Weight', 'Density', 'Float Or Sink', 'Equations of Motion']
def walk(shapes):
    for sh in shapes:
        if sh.shape_type == 6: yield from walk(sh.shapes)
        elif sh.has_text_frame and sh.text_frame.text.strip(): yield sh.text_frame.text.replace('\n', ' ')
old = []; used = []
for name in prev:
    f = ROOT / 'reference' / f'{name}.pptx'
    if not f.exists(): f = ROOT / 'out' / name / f'{name}.pptx'
    if not f.exists(): continue
    used.append(name)
    old += [t_ for t_ in walk(Presentation(str(f)).slides[0].shapes) if len(t_) > 12]
ok(len(used) >= 3 and 'Resultant Forces' in used, f'compared with earlier Do Nows from: {", ".join(used)}')
worst = max((difflib.SequenceMatcher(None, n1[f'd{i}_q'].lower(), b.lower()).ratio(), n1[f'd{i}_q'], b) for i in range(6) for b in old)
print('closest pair:', worst[1][:70], '<>', worst[2][:70])
ok(worst[0] < 0.7, f'no Do Now question repeats one from the last 10A lessons: closest match {worst[0]:.2f}')

# ---- Hook (3) ----
n3 = named(3)
ok('5 km/s' in n3['slide_title'] and 'engine off' in n3['slide_title'] and n3['h1_k'] == 'B' and 'keeps moving at 5 km/s in a straight line' in n3['h1_t'], 'Hook: a probe in deep space; B (keeps moving at 5 km/s in a straight line) is the answer')
ok('A' == n3['h0_k'] and 'slows down and stops' in n3['h0_t'] and 'speeds up' in n3['h2_t'], 'Hook: A (slows and stops) and C (speeds up) are the two traps')
ok('B' in notes(3).split('ANSWER, FOR YOU:')[1][:6], 'the Hook notes give B as the answer')

# ---- I Do 1 (4): the animation's numbers, and the worked example ----
media = (ROOT / 'build' / 'media' / 'when-the-resultant-is-zero-media.py').read_text()
ok("{'name': 'Weight', 'dir': 'down', 'n': 100}" in media and "{'name': 'Reaction force', 'dir': 'up', 'n': 100}" in media and "'name': 'Push', 'dir': 'right', 'n': 60" in media and "'name': 'Friction', 'dir': 'left', 'n': 60" in media and 'V0, PX_PER_MS = 5.0' in media, 'the animation: weight 100 N, reaction 100 N, push 60 N, friction 60 N, a steady 5 m/s')
ok(100 - 100 == 0 and 60 - 60 == 0 and 0 - 60 == -60 and '60 − 60 = 0 N' in media and '0 − 60 = 60 N to the left' in media, 'the animation: 100 − 100 = 0, 60 − 60 = 0 (steady), 0 − 60 = 60 N to the left (the speed falls)')
n4 = named(4)
ok('unless a resultant force acts on it' in n4['fb_t'] and 'constant speed in a straight line' in n4['fb_t'] and 'at rest' in n4['fb_t'], 'I Do 1 states the first law, in the scheme\'s words (at rest, or constant speed in a straight line, unless a resultant force)')
ok('push is 60 N' in n4['wx_t'] and 'resultant is 0 N' in n4['wx_t'] and 'friction equals the push: 60 N' in n4['wx_t'], 'I Do 1 worked example matches the animation: push 60 N, resultant 0 N, friction 60 N')

# ---- I Do 2 (5): zero resultant against no forces ----
n5 = named(5)
ok('ZERO' in n5['lh_t'] and 'cancel' in n5['lh_t'] and 'NO FORCES AT ALL' in n5['rh_t'], 'I Do 2: left is "the resultant is zero: the forces cancel", right is "no forces at all"')
ok('Forces ARE acting' in n5['lf_t'] and 'No forces' in n5['rc_t'] and 'probe' in n5['rc_t'], 'I Do 2: the left card says forces ARE acting; the right card is the probe with no forces')
ok(all(k in n5 for k in ('lr0_h', 'lr1_h', 'lr2_h')) and 'book' in n5['lr0_h'] and 'steady' in n5['lr1_h'] and 'steady' in n5['lr2_h'], 'I Do 2: the three zero-resultant examples (book at rest, car and skydiver at a steady speed)')
ok('resultant' in n5['bn'] and 'same result' in n5['bn'], 'I Do 2 banner: two reasons, the same result, about the resultant')
ok('PROBE' not in n5['lh_t'] and 'book' not in n5['rc_t'].lower(), 'I Do 2: no example is on the wrong side')

# ---- We Do (6) ----
n6 = named(6)
ok(70 - 45 == 25 and '25 N, forwards' in n6['wd1_a'] and 'increases' in n6['wd1_a'], 'We Do row 2: the friction stays 45 N, 70 − 45 = 25 N forwards, the speed increases')
ok('1800 N' in n6['wd0_q'] and '1800 N' in n6['wd0_a'] and 'driving force' in n6['wd0_a'], 'We Do row 1: steady speed, so the resistive force is the driving force, 1800 N')
ok('0 N' in n6['wd2_a'] and 'same speed in a straight line' in n6['wd2_a'], 'We Do row 3: no forces, 0 N, the same speed in a straight line')
ok(all(n6[f'wd{i}_q'].count('____') >= 2 for i in range(3)), 'We Do: each row has at least two blanks to finish (Finish this one)')

# ---- Cold Call (7) and Plenary (10) ----
n7 = named(7)
ok(150 - (90 + 35) == 25 and n7['c4_a'].startswith('25 N to the right'), 'Cold Call Q5: 150 − (90 + 35) = 25 N to the right')
ok(R(98, 10) * 2 == R(196, 10) and n7['c5_a'].startswith('19.6 m/s'), 'Cold Call Q6: 9.8 × 2 = 19.6 m/s')
ok('unless a resultant force acts on it' in n7['c0_a'] and 'no forces on it' in n7['c2_q'] and 'resultant is 0 N' in n7['c2_a'], 'Cold Call: Q1 is the law; Q3 is the collapse ("no forces on it") and its answer')
n10 = named(10)
ok([n10[f'p{i}_v'] for i in range(5)] == ['TRUE', 'FALSE', 'FALSE', 'TRUE', 'FALSE'], 'Plenary: TRUE, FALSE, FALSE, TRUE, FALSE (each FALSE is a misconception from the lesson; Q4 is the applied item)')
ok('no forces on it' in n10['p2_q'] and 'slows down and stops' in n10['p4_q'] and 'is needed to keep' in n10['p1_q'], 'Plenary: the three FALSE statements are the collapse, the deep-space trap and "motion needs a force"')
ok(45 - 45 == 0 and 'friction is 45 N' in n10['p3_q'], 'Plenary Q4: a steady speed, so the friction equals the pull, 45 N')
ok('Next lesson' in n10['pl_next'] and 'NOT zero' in n10['pl_next'] and 'NOT WRITTEN DOWN' in notes(10), 'the closing line points at the next idea, and the notes say that this is an assumption')

# ---- the brief's AVOID: "no resultant force" must never collapse into "no forces" ----
LEGIT = re.compile(r'probe|space|not mean|not the same|nothing pushes|or no forces|student says|ball|plenary|false', re.I)
hits = []
for i in range(1, 11):
    for nm, tx in named(i).items():
        for m in re.finditer(r'no forces?\b', tx, re.I):
            hits.append((i, nm))
            ctx = tx + ' ' + (n10.get('pl_next', '') if i == 10 else '')
            ok(bool(LEGIT.search(tx)) or (i == 10 and nm == 'p2_q') or (i == 7 and nm == 'c2_q') or (i == 2 and nm == 'o2_t'), f'slide {i} {nm}: "no force(s)" appears only where it is meant (deep space, a quoted mistake, or "does not mean"): {tx[:70]}')
ok(any(i == 5 for i, _ in hits) and any(i == 9 for i, _ in hits), 'both I Do 2 and the Mark slide say "no forces", each in its right place')
ok('does not mean there are no forces' in named(9)['mk_card_t'] and 'resultant force' in named(2)['obj_banner'] and 'forces can still be there' in named(2)['obj_banner'], 'the banner and the Mark card both say the forces can still be there / a resultant of 0 N is not no forces')
banned = re.compile(r'angle|diagonal|degree|°|component|vector|pythag|trigonom|parallelogram', re.I)
for i in range(1, 11):
    ok(not banned.search(text(i)), f'slide {i}: no forces at angles')

# ---- the worksheet and its answers ----
js = subprocess.check_output(['node', '-e', "const a=require('./build/when-the-resultant-is-zero-answers.js');console.log(JSON.stringify({a:a,D:a.D,R:a.R}))"], cwd=ROOT).decode()
data = json.loads(js); A = {k: v for k, v in data['a']}; D = data['D']; RR = data['R']
ok(len(A) == 14 and all(A[str(i)] for i in range(1, 15)), 'fourteen worksheet answers, one for every question, in order')
ok(D['w']['push'] == D['w']['friction'] == 30, 'Bronze worked example: a steady speed, so the friction equals the 30 N push')
ok(D['q3']['weight'] - D['q3']['reaction'] == 0 == RR['q3'] and 'Resultant: 12 − 12 = 0 N' in A['3'], 'Q3: the book\'s weight and reaction force are both 12 N, so the resultant is 0 N')
ok(sum(D['q4']['right']) - sum(D['q4']['left']) == 0 == RR['q4'] and '25 + 20 = 45' in A['4'] and '45 − 45 = 0' in A['4'], 'Q4: 45 − (25 + 20) = 0 N, so the speed stays the same')
ok(D['q5']['driving'] == 900 == RR['q5'] and A['5'].startswith('900 N'), 'Q5: a steady speed, so the resistive force is 900 N')
ok(D['q6']['right'][0] - D['q6']['left'][0] == 50 == RR['q6X'] and '200 − 150 = 50' in A['6'], 'Q6: 200 = 150 + X, so X = 50 N')
ok(D['sw']['weight'] == 750, 'Silver worked example: a skydiver at a steady speed, weight 750 N, so the air resistance is 750 N')
ok(D['q10']['after'] == 8 and D['q10']['after'] > D['q10']['before'] and '8 km/s' in A['10'], 'Q10: the probe keeps the NEW speed, 8 km/s')
ok(D['q11']['mass'] * G == R(4900) == RR['q11'] and '500 × 9.8 = 4900' in A['11'], 'Q11: tension = weight = 500 × 9.8 = 4900 N')
ok(D['q13']['new'] - D['q13']['steady'] == 25 == RR['q13'] and '85 − 60 = 25' in A['13'] and 'speeds up' in A['13'], 'Q13: the push goes from 60 N to 85 N, so the resultant is 25 N forwards and the crate speeds up')
ok('not because there are none' in A['9'] and 'cancel' in A['9'] and 'cancel' in A['14'] and 'no forces at all' in A['14'].lower(), 'Q9 and Q14 answers keep "the forces cancel" and "no forces at all" apart')
ok(all(w in A['8'] for w in ('no resultant force', 'no change in motion', 'cancel')) and 'resultant' in A['2'] and 'constant speed' in A['2'], 'Q2 and Q8 answers use the first law')

# the free-body diagrams the worksheet prints must show exactly the numbers the answers use
ws = ast.literal_eval(re.search(r'WS = (\{.*?\n\})', media, re.S).group(1))
f4, lab4, _ = ws['wtrz-ws-q4']; f6, lab6, _ = ws['wtrz-ws-q6']
ok(sorted(f['n'] for f in f4 if f['dir'] == 'right') == sorted(D['q4']['right']) and sorted(f['n'] for f in f4 if f['dir'] == 'left') == sorted(D['q4']['left']), f'the Q4 diagram ({lab4}) draws exactly the forces in the answers')
ok([f['n'] for f in f6 if f['dir'] == 'right'] == D['q6']['right'] and [f['n'] for f in f6 if f['dir'] == 'left' and 'label' not in f] == D['q6']['left'] and [f['n'] for f in f6 if 'label' in f] == [RR['q6X']], f'the Q6 diagram ({lab6}) draws thrust 200, water resistance 150 and the unknown X to scale (50)')

z = zipfile.ZipFile(OUT / f'{LESSON} worksheet.docx')
body = re.sub(r'\s+', ' ', re.sub(r'<[^>]+>', ' ', z.read('word/document.xml').decode('utf-8')))
for needle in ['push is 30 N', '12 N rests on a table', 'driving force is 900 N', '750 N falls at a steady speed', 'Right: 200 N', 'speeds up from 3 km/s to 8 km/s', 'mass 500 kg', 'g = 9.8 N/kg', 'increased to 85 N', 'force of 60 N', 'resultant force is zero', 'there are no forces']:
    ok(needle in body, f'the worksheet prints "{needle}"')
ok(not banned.search(body), 'the worksheet has no forces at angles')
media_parts = [n for n in z.namelist() if n.startswith('word/media/') and not n.endswith('/')]
ok(len(media_parts) == 3 and all(n.endswith('.png') for n in media_parts), f'the worksheet has its {len(media_parts)} images (2 diagrams and the upside-down answers) as .png parts (not ".undefined")')
wsrc = (ROOT / 'build' / 'when-the-resultant-is-zero-worksheet.js').read_text()
for i in range(1, 15): ok(re.search(rf"q\('{i}',", wsrc), f'worksheet has question {i}')
ok("q('8'" in wsrc and 'Why does that step work?' in wsrc and 'Why does a steady speed mean the resultant force is zero' in wsrc, 'the "why does that step work?" prompt (Q8) is on why a steady speed means a zero resultant')
ok("q('12'" in wsrc and 'Where would you meet this idea outside the lesson' in wsrc, 'the "where would you meet this outside the lesson?" question is Q12')
ok((OUT / f'{LESSON} worksheet.docx').exists(), 'the worksheet exists')

# ---- the game: the deck points at a file, so the file must exist ----
ok(f'{LESSON} game' in text(8), 'slide 8 names "When The Resultant Is Zero game"')
game = OUT / f'{LESSON} game.html'
ok(game.exists() and game.stat().st_size > 20000, f'the game file the slide points at exists ({game.stat().st_size // 1024} KB)')
html = game.read_text(encoding='utf-8')
ok(html == (ROOT / 'build' / 'when-the-resultant-is-zero-game.html').read_text(encoding='utf-8'), 'the delivered game is exactly the source file')
ok('9.8' in html and 'Zero Or Not' in html, 'the game uses g = 9.8 N/kg')
print(f'\n{checks} checks passed')
