#!/usr/bin/env python3
"""
Checks everything numeric (and every claim that can be checked) in Technology And Society: the
deck, the worksheet and its answers. Values are read BACK OUT of the finished files and recomputed
here, independently of the JavaScript that wrote them.

    python3 build/technology-and-society-check.py

It also fails if the deck points at a game file that does not exist, and if the deck's phases do
not add to 50 (ten slides, You Do 14, Mark 3: TEMPLATE.md).
"""
import datetime, json, pathlib, re, subprocess
from decimal import Decimal
from fractions import Fraction
from pptx import Presentation

ROOT = pathlib.Path(__file__).parent.parent
OUT = ROOT / 'out' / 'Technology And Society'
DECK = OUT / 'Technology And Society.pptx'
checks = 0
def ok(cond, msg):
    global checks
    assert cond, f'FAIL: {msg}'
    checks += 1
    print('ok:', msg)

prs = Presentation(str(DECK))
slides = list(prs.slides)
def walk(shapes):
    # the animation step wraps shapes in groups, so look inside them too
    for sh in shapes:
        if sh.shape_type == 6: yield from walk(sh.shapes)
        elif sh.has_text_frame and sh.text_frame.text.strip(): yield sh.text_frame.text.replace('\n', ' ')
def text(i): return ' | '.join(walk(slides[i - 1].shapes))
def notes(i): return slides[i - 1].notes_slide.notes_text_frame.text

# ---- the shape: TEMPLATE.md, ten slides, You Do 14, Mark 3, 50 minutes ----
pills = []
for i in range(1, len(slides) + 1):
    m = re.search(r'([A-Z][A-Z ]+) · (\d+) MIN', text(i)); pills.append((m.group(1), int(m.group(2))) if m else None)
ok(len(slides) == 10 and all(pills), 'ten slides, each with a phase pill')
ok([p[0] for p in pills] == ['DO NOW', 'TODAY', 'HOOK', 'I DO', 'I DO', 'WE DO', 'COLD CALL', 'YOU DO', 'MARK', 'PLENARY'], 'phases in the TEMPLATE.md order')
ok([p[1] for p in pills] == [10, 1, 2, 3, 3, 5, 6, 14, 3, 3] and sum(p[1] for p in pills) == 50, 'minutes 10, 1, 2, 3, 3, 5, 6, 14, 3, 3 add to 50')
ok(all(notes(i).strip() for i in range(1, 11)), 'speaker notes on every slide')

# ---- the date ----
d = re.search(r'(Monday|Tuesday|Wednesday|Thursday|Friday|Saturday|Sunday) (\d+) (\w+) (\d{4})', text(1))
dt = datetime.datetime.strptime(f'{d.group(2)} {d.group(3)} {d.group(4)}', '%d %B %Y')
ok(dt.strftime('%A') == d.group(1), f'the date on slide 1 is a real {d.group(1)}: {d.group(0)}')

# ---- Do Now (slide 1) ----
s1 = text(1)
nums = [Fraction(x) for x in re.search(r'Four results: (\d+), (\d+), (\d+) and (\d+)', s1).groups()]
mean = sum(nums) / len(nums); rng = max(nums) - min(nums)
ok(mean == 25 and rng == 6, f'Do Now Q1: {[int(n) for n in nums]} has mean {mean} and range {rng}')
ok('Mean 25 (100 ÷ 4). Range 6 (28 − 22).' in s1 and sum(nums) == 100, 'Do Now Q1: the printed answer says 25, 100 ÷ 4, 6, 28 − 22 and the total is 100')
ok('A bar chart' in s1 and 'fertiliser' in s1 and 'Independent: the fertiliser. Dependent: the plant height.' in s1, 'Do Now Q2 to Q4 present with their answers')
ok('law says what happens' in s1.lower() and 'theory explains why' in s1.lower(), 'Do Now Q5 (last term): a law says what, a theory says why')
ok(sum(1 for q in ['Four results', 'best graph', 'fertiliser makes', 'same fertiliser', 'law and a scientific theory', 'invention'] if q in s1) == 6, 'Do Now has six questions: 2 last lesson, 2 earlier in the unit, 1 last term, 1 preview')

# ---- Cold Call (slide 7): the accuracy question and the mean ----
s7 = text(7)
reads = [Decimal(x) for x in re.search(r'reads ([\d.]+) °C, ([\d.]+) °C and ([\d.]+) °C', s7).groups()]
true = Decimal(re.search(r'really (\d+) °C', s7).group(1))
tm = sum(reads) / len(reads); off = true - tm; rg = max(reads) - min(reads)
ok(tm == Decimal('21.0') and off == Decimal('4.0') and rg == Decimal('0.2'), f'Cold Call Q2: readings {[str(r) for r in reads]} have mean {tm}, are {off} °C below the true {true} °C, and range {rg}')
ok(rg < Decimal('0.5') and off > 1 and 'Precise, not accurate. Close together, but 4 °C too low.' in s7, 'Cold Call Q2: close together (precise) and far from the true value (not accurate), and the printed answer says 4 °C')
m4 = [Fraction(x) for x in re.search(r'mean of (\d+), (\d+), (\d+) and (\d+)', s7).groups()]
ok(sum(m4) / 4 == 7 and '7. (4 + 6 + 8 + 10) ÷ 4 = 28 ÷ 4.' in s7 and sum(m4) == 28, 'Cold Call Q4: the mean of 4, 6, 8 and 10 is 7 (28 ÷ 4)')

# ---- the worksheet: its worked example and half-worked example (read from the answers file) ----
ans = json.loads(subprocess.check_output(['node', '-e', "const a=require('./build/technology-and-society-answers.js');console.log(JSON.stringify({a:a,mean:a.MEAN,scores:a.SCORES}))"], cwd=ROOT).decode())
A = {k: v for k, v in ans['a']}
sc = [Fraction(x) for x in ans['scores']]
ok(sc == [3, 4, 2, 3] and sum(sc) / len(sc) == 3 and ans['mean'] == 3, 'worksheet Q6: scores 3, 4, 2, 3 have mean (3 + 4 + 2 + 3) ÷ 4 = 12 ÷ 4 = 3')
ok(sum(sc) / len(sc) > 2 and 'more than 2' in A['6'] and '= 12 ÷ 4 = 3' in A['6'], 'worksheet Q6: 3 is more than "2 or less", so the design does not yet work; the printed answer says so')
ok(Fraction(2 + 1 + 3, 3) == 2, 'worksheet worked example: scores 2, 1 and 3 give a mean of (2 + 1 + 3) ÷ 3 = 6 ÷ 3 = 2, so the bag works (2 or less)')
ws = (ROOT / 'build' / 'technology-and-society-worksheet.js').read_text()
ok('(2 + 1 + 3) ÷ 3 = 6 ÷ 3 = 2' in ws and '(3 + 4 + 2 + 3) ÷ ___ = ___ ÷ ___ = ___' in ws, 'the worksheet prints those two calculations')
ok(len(A) == 13 and all(A[str(i)] for i in range(1, 14)), 'thirteen worksheet answers, one for every question, in order')
for i in range(1, 14): ok(re.search(rf"q\('{i}',", ws), f'worksheet has question {i}')

# ---- the same numbers in the deck and the worksheet agree (5 kg, 100 m, 3 times) ----
s5 = text(5)
ok('Pull 5 kg for 100 m. Three students, three times each.' in s5 and 'loading the bag with 5 kg and pulling it 100 m' in ws, 'I Do 2 and the worksheet use the same test: 5 kg, 100 m, repeated 3 times')

# ---- dated facts on the slides (each checked with a web search when the deck was written) ----
ok('1980' in text(4) and '1980' in notes(3) and '8 May 1980' in notes(3), 'smallpox: declared eradicated by the World Health Assembly on 8 May 1980 (WHO)')
ok('1987' in text(6) and 'Montreal Protocol' in notes(6) and '2066' in notes(6), 'CFCs and the ozone layer: the Montreal Protocol was 1987, and Antarctic ozone is projected to recover by about 2066 (WMO/UNEP 2022)')

# ---- the game: the deck points at a file, so the file must exist ----
ok('Technology And Society game' in text(8), 'slide 8 names "Technology And Society game"')
game = OUT / 'Technology And Society game.html'
ok(game.exists() and game.stat().st_size > 20000, f'the game file the slide points at exists ({game.stat().st_size // 1024} KB)')
html = game.read_text(encoding='utf-8')
data = json.load((ROOT / 'build' / 'technology-and-society.data.json').open(encoding='utf-8')); data.pop('note', None)
ok(json.dumps(data, ensure_ascii=False) in html, 'the game contains exactly the briefs in technology-and-society.data.json')
ok((OUT / 'Technology And Society worksheet.docx').exists(), 'the worksheet exists')
print(f'\n{checks} checks passed')
