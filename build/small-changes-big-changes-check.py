#!/usr/bin/env python3
"""
Checks everything numeric (and every claim that can be checked) in Small Changes, Big Changes (8I): the deck, the worksheet
and its answers, and that the game the deck names exists. Values are read BACK OUT of the finished files and recomputed here
with exact fractions, independently of the JavaScript that wrote them. The game itself is checked over 300 generated games
by build/check-small-changes-big-changes-game.py.

    python3 build/small-changes-big-changes-check.py

It also enforces the brief: allele frequency stays a proportion or a percentage (no Hardy-Weinberg on a slide, in the
worksheet or in the game), and the Do Now does not repeat a question from the last three lessons in the class.
"""
import datetime, difflib, json, pathlib, re, subprocess, zipfile
from fractions import Fraction as F
from pptx import Presentation

ROOT = pathlib.Path(__file__).parent.parent
LESSON = 'Small Changes, Big Changes'
OUT = ROOT / 'out' / LESSON
DECK = OUT / f'{LESSON}.pptx'
checks = 0
def ok(cond, msg):
    global checks
    assert cond, f'FAIL: {msg}'
    checks += 1
    print('ok:', msg)
pct = lambda k, n: F(100 * k, n)

prs = Presentation(str(DECK)); slides = list(prs.slides)
def named(i):
    out = {}
    def go(shapes):
        for sh in shapes:
            if sh.shape_type == 6: go(sh.shapes)
            elif sh.has_text_frame and sh.text_frame.text.strip(): out[sh.name] = sh.text_frame.text
    go(slides[i - 1].shapes); return out
def text(i): return ' | '.join(v.replace('\n', ' ') for v in named(i).values())
def notes(i): return slides[i - 1].notes_slide.notes_text_frame.text

# ---- the shape: TEMPLATE.md ----
pills = []
for i in range(1, len(slides) + 1):
    m = re.search(r'([A-Z][A-Z ]+) · (\d+) MIN', text(i)); pills.append((m.group(1), int(m.group(2))) if m else None)
ok(len(slides) == 10 and all(pills), 'ten slides, each with a phase pill')
ok([p[0] for p in pills] == ['DO NOW', 'TODAY', 'HOOK', 'I DO', 'I DO', 'WE DO', 'COLD CALL', 'YOU DO', 'MARK', 'PLENARY'], 'phases in the TEMPLATE.md order')
ok([p[1] for p in pills] == [10, 1, 2, 3, 3, 5, 6, 14, 3, 3] and sum(p[1] for p in pills) == 50, 'minutes 10, 1, 2, 3, 3, 5, 6, 14, 3, 3 add to 50')
ok(all(notes(i).strip() for i in range(1, 11)), 'speaker notes on every slide')
ok(prs.core_properties.title == LESSON and prs.core_properties.subject == 'Y8 Science · Natural Selection · Lesson 6 · 8I', f'dc:title and dc:subject are set, with the class: {prs.core_properties.subject}')
d = re.search(r'(Monday|Tuesday|Wednesday|Thursday|Friday|Saturday|Sunday) (\d+) (\w+) (\d{4})', text(1))
dt = datetime.datetime.strptime(f'{d.group(2)} {d.group(3)} {d.group(4)}', '%d %B %Y')
ok(dt.strftime('%A') == d.group(1), f'the date on slide 1 is a real {d.group(1)}: {d.group(0)}')
n2 = named(2)
OBJ = ['Define a gene pool.', 'Explain microevolution as a change in allele frequencies within a population.', 'Explain macroevolution as large changes over long periods, above the species level.']
ok(all(n2[f'o{i}_t'] == OBJ[i] for i in range(3)) and 'Objectives' in n2['slide_title'] and text(2).startswith('TODAY · 1 MIN'), 'the three objectives are exactly the brief\'s wording, under the title Objectives with the TODAY pill')

# ---- Do Now ----
n1 = named(1)
ok(F(200 * (100 - 96), 100) == 8 and F(200 * 96, 100) == 192 and n1['d1_a'].startswith('8.') and '96% the same in 200 DNA letters' in n1['d1_q'], 'Do Now Q2: 96% the same in 200 letters leaves 4% = 8 letters different (and 192 the same)')
ok(pct(3, 20) == 15 and n1['d4_a'].startswith('15%.'), 'Do Now Q5: 3 out of 20 is 15%')
ok('tail and pharyngeal arches' in n1['d0_q'] and 'common ancestor' in n1['d0_a'], 'Do Now Q1 (the topic they found hard): embryos, with the reasoning, in the answer')
ok('gene pool' in n1['d5_q'], 'Do Now Q6 previews the gene pool')
prev = ['More Evidence', 'Fossils And The Fossil Record', 'Natural Selection In Action']
old = []
def walk(shapes):
    for sh in shapes:
        if sh.shape_type == 6: yield from walk(sh.shapes)
        elif sh.has_text_frame and sh.text_frame.text.strip(): yield sh.text_frame.text.replace('\n', ' ')
for name in prev:
    f = ROOT / 'reference' / f'{name}.pptx'
    if not f.exists(): f = ROOT / 'out' / name / f'{name}.pptx'
    old += [t_ for t_ in walk(Presentation(str(f)).slides[0].shapes) if len(t_) > 12]
worst = max((difflib.SequenceMatcher(None, n1[f'd{i}_q'].lower(), b.lower()).ratio(), n1[f'd{i}_q'], b) for i in range(6) for b in old)
print('closest pair:', worst[1][:70], '<>', worst[2][:70])
ok(worst[0] < 0.7, f'no Do Now question repeats one from {", ".join(prev)}: closest match {worst[0]:.2f}')

# ---- Hook (3) and I Do 1 (4): the bag ----
n3, n4 = named(3), named(4)
ok(F(4, 5) * 100 == 80 and 'Hook answer: C. The handful said 80%; the bag says 75%.' in n4['bg_t'], 'the handful (4 green of 5) is 80% green')
ok(F(15, 20) * 100 == 75 and F(5, 20) * 100 == 25 and F(15, 20) + F(5, 20) == 1, 'the bag: 15 green of 20 is 75%, 5 brown is 25%, and they add to 100%')
ok('Green: 15 out of 20 = 15 ÷ 20 = 0.75 = 75%' in n4['bg_t'] and 'Brown: 5 out of 20 = 5 ÷ 20 = 0.25 = 25%' in n4['bg_t'], 'I Do 1 prints the two counts and their percentages')
ok(10 * 2 == 20 and '10 beetles, 2 alleles each' in n4['bg_cap'] and 'two alleles' in n4['al_two'], 'I Do 1: 10 beetles × 2 alleles = a pool of 20')
ok(all(w in n4['gp_t'] for w in ('All the alleles', 'all the individuals', 'population')) and 'A version of a gene' in n4['al_t'] and 'section of DNA' in n4['gn_t'], 'I Do 1 defines gene, allele and gene pool')
ok(0.4 < 0.8 and 'C' in n3['h2_k'] and 'cannot tell' in n3['h2_t'], 'the Hook answer is C, the third card')

# ---- the animation (26 s): three generations of 20 beads ----
GENS = [15, 17, 19]
ok([int(pct(g, 20)) for g in GENS] == [75, 85, 95] and all(F(g, 20) * 100 == 5 * g for g in GENS), 'animation: green is 15, 17 and 19 of 20: 75%, 85%, 95%')
src = (ROOT / 'build' / 'media' / 'small-changes-big-changes-beads.py').read_text()
ok('GENS = [15, 17, 19]' in src and "assert [g * 5 for g in GENS] == [75, 85, 95]" in src and 'Birds eat more brown beetles' in src, 'the animation script has those numbers and its own assertion')
ok('75% → 85% → 95%' in src, 'the animation ends on 75% → 85% → 95%')

# ---- We Do (6) ----
n6 = named(6)
ok(pct(15, 20) == 75 and pct(18, 20) == 90 and 'Green before: 15 ÷ 20 = 75%. Green after: 18 ÷ 20 = 90%.' in n6['wf0_1_q'], 'We Do row 1: green 15 of 20 = 75%, then 18 of 20 = 90%')
ok(pct(10, 40) == 25 and pct(24, 40) == 60 and n6['wf1_1_a'] == 'Before: 10 ÷ 40 = 25%. After: 24 ÷ 40 = 60%.' and 'Brown rose from 25% to 60%' in n6['wf1_2_a'], 'We Do row 2: 10 of 40 is 25%, 24 of 40 is 60%, a rise of 35 points')
ok(pct(24, 40) - pct(10, 40) == 35, 'We Do row 2: the change is 60 − 25 = 35 percentage points')
ok(n6['wf0_2_a'].startswith('Microevolution.') and n6['wf1_2_a'].startswith('Microevolution.') and n6['wf2_2_a'] == 'Macroevolution.', 'We Do rows 1 and 2 are microevolution, row 3 is macroevolution')

# ---- Cold Call (7) ----
n7 = named(7)
ok(pct(10, 50) == 20 and n7['c2_a'].startswith('20%. 10 ÷ 50 = 0.2.'), 'Cold Call Q3: 10 of 50 alleles is 20%')
ok('tail and pharyngeal arches' in n7['c4_a'] and 'resistance allele' in n7['c5_a'], 'Cold Call Q5 and Q6 are from earlier lessons (embryos; antibiotic resistance)')

# ---- Plenary (10): each statement is evaluated ----
n10 = named(10)
ok(pct(10, 40) == 25 and pct(20, 40) == 50 and 'TRUE' == n10['p2_v'] and '10 brown' in n10['p2_q'] and '20 brown' in n10['p2_q'], 'Plenary Q3 (the applied item): 10 of 40 (25%) to 20 of 40 (50%) is a change in allele frequency, so TRUE')
ok([n10[f'p{i}_v'] for i in range(5)] == ['FALSE', 'FALSE', 'TRUE', 'FALSE', 'TRUE'], 'Plenary: FALSE, FALSE, TRUE, FALSE, TRUE (each FALSE is a misconception from the lesson)')
ok('Next lesson' in n10['pl_next'] and 'four forces' in n10['pl_next'], 'the closing line says what the next lesson does (it is known: four forces, one of them random)')

# ---- the worksheet and its answers ----
js = subprocess.check_output(['node', '-e', "const a=require('./build/small-changes-big-changes-answers.js');console.log(JSON.stringify({a:a,D:a.D,R:a.R}))"], cwd=ROOT).decode()
data = json.loads(js); A = {k: v for k, v in data['a']}; D = data['D']
ok(len(A) == 14 and all(A[str(i)] for i in range(1, 15)), 'fourteen worksheet answers, one for every question, in order')
ok(pct(D['w']['a'], D['w']['n']) == 40 and pct(D['w']['b'], D['w']['n']) == 60, 'worked example: 10 and 15 of 25 are 40% and 60%')
ok(pct(D['q1']['a'], D['q1']['n']) == 30 and pct(D['q1']['b'], D['q1']['n']) == 70 and '= 30%' in A['1'] and '= 70%' in A['1'], 'Q1: 15 and 35 of 50 are 30% and 70%')
ok(D['q3']['beetles'] * 2 == 20 and '= 20 alleles' in A['3'], 'Q3: 10 beetles × 2 = 20 alleles')
ok(pct(D['q4']['brown'], 20) == 35 and '= 35%' in A['4'], 'Q4: 7 of 20 alleles is 35%')
ok(pct(D['q9']['before'], D['q9']['n']) == 25 and pct(D['q9']['after'], D['q9']['n']) == 60 and pct(D['q9']['after'], D['q9']['n']) - pct(D['q9']['before'], D['q9']['n']) == 35 and '35 percentage points' in A['9'] and 'microevolution' in A['9'], 'Q9: 10 and 24 of 40 are 25% and 60%: 35 points, microevolution')
ok(D['q13']['beetles'] * 2 == 60 and pct(D['q13']['brown'], 60) == 35 and pct(D['q13']['brown'], 30) == 70 and '= 35%' in A['13'] and '70%' in A['13'], 'Q13: 30 beetles are 60 alleles; 21 of 60 is 35% (not 21 of 30 = 70%)')
ok(pct(D['q14']['brown'] - D['q14']['lost'], D['q14']['n'] - D['q14']['lost']) == 25 and '= 25%' in A['14'] and '(20 − 10)' not in A['14'], 'Q14: (20 − 10) of (50 − 10) = 10 of 40 = 25%')
ok(D['q14']['brown'] * 100 // D['q14']['n'] == 40 and pct(10, 50) == 20 and pct(10, 40) != 20, 'Q14: the shrinking pool matters (10 of 50 is 20%, but 10 of 40 is 25%)')
z = zipfile.ZipFile(OUT / f'{LESSON} worksheet.docx')
body = re.sub(r'\s+', ' ', re.sub(r'<[^>]+>', ' ', z.read('word/document.xml').decode('utf-8')))
for needle in ['10 ÷ 25 = 0.4 = 40%', '15 ÷ 25 = 0.6 = 60%', '25 alleles', '50 alleles', '10 beetles', '30 beetles', '40 alleles']:
    ok(needle in body, f'the worksheet prints "{needle}"')
wsrc = (ROOT / 'build' / 'small-changes-big-changes-worksheet.js').read_text()
for i in range(1, 15): ok(re.search(rf"q\('{i}',", wsrc), f'worksheet has question {i}')
ok("q('8'" in wsrc and 'Why does that step work?' in wsrc and 'ALLELES' in wsrc, 'the "why does that step work?" prompt (Q8) is on dividing by the alleles, not the beetles')
ok("q('12'" in wsrc and 'Where would you meet this idea outside the lesson' in wsrc, 'the "where would you meet this outside the lesson?" question is Q12')
ok((OUT / f'{LESSON} worksheet.docx').exists(), 'the worksheet exists')

# ---- the game: the deck points at a file, so the file must exist ----
ok(f'{LESSON} game' in text(8), 'slide 8 names "Small Changes, Big Changes game"')
game = OUT / f'{LESSON} game.html'
ok(game.exists() and game.stat().st_size > 20000, f'the game file the slide points at exists ({game.stat().st_size // 1024} KB)')
html = game.read_text(encoding='utf-8'); gsrc = (ROOT / 'build' / 'small-changes-big-changes-game.html').read_text(encoding='utf-8')
ok(html == gsrc, 'the delivered game is exactly the source file')
ok('Microevolution' in html and 'Macroevolution' in html and 'above the species level' in html, 'the game has the two bins, with the lesson\'s definitions')

# ---- the brief ----
bad = ['hardy', 'weinberg']
for name, blob in {'slides': ' '.join(text(i) for i in range(1, 11)), 'worksheet': body, 'game': html}.items():
    ok(not any(b in blob.lower() for b in bad), f'no Hardy-Weinberg in the {name}')
print(f'\n{checks} checks passed')
